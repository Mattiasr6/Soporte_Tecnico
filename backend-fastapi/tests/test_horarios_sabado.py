"""Tests de horarios especiales de sábado por fecha.

Cupo: mañana 2, mediodía 1, tarde 1-2. Solo sábados (YYYY-MM-DD).
"""

import io
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.models.usuario import Usuario  # noqa: F401  (registra el modelo)

client = TestClient(app)

UID_TEC = 2
UID_JEFE = 8
EMAILS = {
    2: "diego.orihuela@upds.edu.bo",
    8: "josue.huayllas@upds.edu.bo",
}
_tokens: dict[int, str] = {}


def h(uid: int) -> dict[str, str]:
    if uid not in _tokens:
        r = client.post(
            "/api/auth/login",
            json={"email": EMAILS[uid], "password": os.environ["SEED_PASSWORD"]},
        )
        assert r.status_code == 200, r.text
        _tokens[uid] = r.json()["token"]
    return {"Authorization": f"Bearer {_tokens[uid]}"}


DATA_DIR = Path(__file__).resolve().parent.parent / "data"
EQUIPO_FILE = DATA_DIR / "equipo_auxiliares.json"
HORARIOS_FILE = DATA_DIR / "horarios_auxiliares.json"
SABADOS_FILE = DATA_DIR / "horarios_sabados.json"


@pytest.fixture
def archivos_data():
    respaldo = {
        p: p.read_bytes() if p.exists() else None
        for p in (EQUIPO_FILE, HORARIOS_FILE, SABADOS_FILE)
    }
    client.put("/api/laboratorios/equipo", json={"auxiliares": []}, headers=h(UID_JEFE))
    if SABADOS_FILE.exists():
        SABADOS_FILE.unlink()
    yield
    for ruta, contenido in respaldo.items():
        if contenido is None:
            if ruta.exists():
                ruta.unlink()
        else:
            ruta.write_bytes(contenido)


AUX = "TEST-SAB-AUX-"
SABADO_OK = "2099-01-03"  # sábado
LUNES = "2099-01-05"  # lunes


def _agregar_aux(nombre: str) -> None:
    r = client.post(
        "/api/laboratorios/equipo", json={"nombre": nombre}, headers=h(UID_JEFE)
    )
    assert r.status_code in (200, 400), r.text


def _body(*nombres: str) -> dict[str, dict[str, object]]:
    return {
        "mañana": {"inicio": "08:00", "fin": "12:00", "auxiliares": list(nombres[0:2])},
        "mediodia": {"inicio": "12:00", "fin": "16:00", "auxiliares": list(nombres[2:3])},
        "tarde": {"inicio": "14:30", "fin": "18:30", "auxiliares": list(nombres[3:])},
    }


def test_put_y_get_del_mes(archivos_data):
    for i in range(1, 6):
        _agregar_aux(f"{AUX}{i}")
    r = client.put(
        f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
        json=_body(*[f"{AUX}{i}" for i in range(1, 6)]),
        headers=h(UID_JEFE),
    )
    assert r.status_code == 200, r.text
    assert r.json()["mañana"]["auxiliares"] == [f"{AUX}1", f"{AUX}2"]
    r = client.get(
        "/api/laboratorios/horarios-sabado",
        params={"mes": 1, "anio": 2099},
        headers=h(UID_TEC),
    )
    assert r.status_code == 200, r.text
    data = r.json()
    fechas = [s["fecha"] for s in data["sabados"]]
    assert SABADO_OK in fechas
    assert LUNES not in fechas
    fila = next(s for s in data["sabados"] if s["fecha"] == SABADO_OK)
    assert fila["bloques"]["tarde"]["auxiliares"] == [f"{AUX}4", f"{AUX}5"]
    assert data["conteo"] == {f"{AUX}{i}": 1 for i in range(1, 6)}
    assert data["defaults"]["tarde"] == {"inicio": "14:30", "fin": "18:30"}


def test_tarde_con_uno_vale_y_con_tres_no(archivos_data):
    for i in range(1, 6):
        _agregar_aux(f"{AUX}{i}")
    ok = _body(*[f"{AUX}{i}" for i in range(1, 5)])
    ok["tarde"]["auxiliares"] = [f"{AUX}4"]
    r = client.put(
        f"/api/laboratorios/horarios-sabado/{SABADO_OK}", json=ok, headers=h(UID_JEFE)
    )
    assert r.status_code == 200, r.text
    mala = _body(*[f"{AUX}{i}" for i in range(1, 6)])
    mala["tarde"]["auxiliares"] = [f"{AUX}4", f"{AUX}5", f"{AUX}1"]
    r = client.put(
        f"/api/laboratorios/horarios-sabado/{SABADO_OK}", json=mala, headers=h(UID_JEFE)
    )
    assert r.status_code == 400


def test_rechazos(archivos_data):
    _agregar_aux(f"{AUX}1")
    base = _body(f"{AUX}1", f"{AUX}1", f"{AUX}1", f"{AUX}1")
    # 403 técnico
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=base,
            headers=h(UID_TEC),
        ).status_code
        == 403
    )
    # lunes no es sábado
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{LUNES}",
            json=base,
            headers=h(UID_JEFE),
        ).status_code
        == 400
    )
    # fecha malformada
    assert (
        client.put(
            "/api/laboratorios/horarios-sabado/ayer",
            json=base,
            headers=h(UID_JEFE),
        ).status_code
        in (400, 404, 422)
    )
    # mañana con 1 solo
    poco = _body(f"{AUX}1", f"{AUX}1", f"{AUX}1", f"{AUX}1")
    poco["mañana"]["auxiliares"] = [f"{AUX}1"]
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=poco,
            headers=h(UID_JEFE),
        ).status_code
        == 400
    )
    # fantasma fuera de nómina
    fantasma = _body(f"{AUX}1", f"{AUX}1", f"{AUX}1", f"{AUX}1")
    fantasma["tarde"]["auxiliares"] = ["No Existe Jamas"]
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=fantasma,
            headers=h(UID_JEFE),
        ).status_code
        == 400
    )
    # repetido en dos turnos
    for i in range(2, 6):
        _agregar_aux(f"{AUX}{i}")
    doble = _body(*[f"{AUX}{i}" for i in range(1, 6)])
    doble["mediodia"]["auxiliares"] = [f"{AUX}1"]
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=doble,
            headers=h(UID_JEFE),
        ).status_code
        == 400
    )


def test_export_xlsx_sabado_y_semanal(archivos_data):
    from openpyxl import load_workbook

    for i in range(1, 6):
        _agregar_aux(f"{AUX}{i}")
    body = _body(*[f"{AUX}{i}" for i in range(1, 6)])
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=body,
            headers=h(UID_JEFE),
        ).status_code
        == 200
    )
    r = client.get(
        "/api/laboratorios/horarios/export.xlsx",
        params={"tipo": "sabado", "mes": 1, "anio": 2099},
        headers=h(UID_TEC),
    )
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"
    assert "sabados_2099-01.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(filename=io.BytesIO(r.content)).active
    assert ws is not None
    assert ws["A1"].value == "HORARIOS TURNO SABADO 01-2099"
    assert [c.value for c in ws[2]] == ["NOMBRE", "SABADO", "TURNO", "INICIO", "FIN"]
    nombres = [fila[0].value for fila in ws.iter_rows(min_row=3)]
    assert f"{AUX}1" in nombres
    libres = [fila for fila in ws.iter_rows(min_row=3) if fila[2].value == "Libre"]
    assert not libres  # el equipo de prueba son solo los 5 asignados
    r = client.get(
        "/api/laboratorios/horarios/export.xlsx",
        params={"tipo": "semanal"},
        headers=h(UID_TEC),
    )
    assert r.status_code == 200, r.text
    assert (
        client.get(
            "/api/laboratorios/horarios/export.xlsx",
            params={"tipo": "word"},
            headers=h(UID_TEC),
        ).status_code
        == 400
    )
    r = client.get(
        "/api/laboratorios/horarios/export.pdf",
        params={"tipo": "sabado", "mes": 1, "anio": 2099},
        headers=h(UID_TEC),
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"
    assert "sabados_2099-01.pdf" in r.headers["content-disposition"]


def test_delete_limpia_la_fecha(archivos_data):
    for i in range(1, 5):
        _agregar_aux(f"{AUX}{i}")
    body = _body(*[f"{AUX}{i}" for i in range(1, 5)])
    assert (
        client.put(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}",
            json=body,
            headers=h(UID_JEFE),
        ).status_code
        == 200
    )
    assert (
        client.delete(
            f"/api/laboratorios/horarios-sabado/{SABADO_OK}", headers=h(UID_TEC)
        ).status_code
        == 403
    )
    r = client.delete(
        f"/api/laboratorios/horarios-sabado/{SABADO_OK}", headers=h(UID_JEFE)
    )
    assert r.status_code == 204, r.text
    data = client.get(
        "/api/laboratorios/horarios-sabado",
        params={"mes": 1, "anio": 2099},
        headers=h(UID_JEFE),
    ).json()
    fila = next(s for s in data["sabados"] if s["fecha"] == SABADO_OK)
    assert fila["bloques"] is None
    assert data["conteo"] == {}
