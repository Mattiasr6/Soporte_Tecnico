"""XLSX/PDF export of the auxiliar schedules (Django `auxiliares/horarios/export.*`).

Same content and look as the Django export (built by the FastAPI
`/api/laboratorios/horarios/export.*` routes), read from the horarios tables:

- semanal: every active auxiliar/encargado with the shift in force today
  (fn_turno_vigente, else the usual shift) and that shift's hours;
- sabado: every assignment of the Saturdays of a month with its hours that day
  (own hours, else `horarios_turno`), then a "Libre" row for each active team
  member without a Saturday that month.
openpyxl and reportlab are already backend dependencies.
"""

import io
from datetime import time
from typing import Any

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.services.asignacion import sabados
from app.services.asignacion.sql import rows

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOMBRE_TURNO = {"M": "Mañana", "MD": "Mediodía", "T": "Tarde", "N": "Noche"}
_ORDEN = {"M": 0, "MD": 1, "T": 2, "N": 3}

Fila = list[Any]


def _hora(valor: time | None) -> str | None:
    return valor.strftime("%H:%M") if valor else None


def _equipo(db: Session) -> list[dict[str, Any]]:
    sql = text(
        """
        select id, nombre_completo,
               coalesce(horarios.fn_turno_vigente(id), turno_habitual) as turno
          from horarios.perfiles
         where activo and rol in ('auxiliar', 'encargado')
         order by nombre_completo
        """
    )
    return [dict(r) for r in rows(db, sql)]


def semanal(db: Session) -> tuple[str, list[str], list[Fila], str]:
    base = sabados.horarios_base(db)
    equipo = sorted(
        _equipo(db), key=lambda m: (_ORDEN.get(m["turno"], 9), m["nombre_completo"])
    )
    filas: list[Fila] = []
    for m in equipo:
        if m["turno"]:
            inicio, fin = base.get(m["turno"], (None, None))
            filas.append(
                [
                    m["nombre_completo"],
                    NOMBRE_TURNO[m["turno"]],
                    _hora(inicio),
                    _hora(fin),
                ]
            )
        else:
            filas.append([m["nombre_completo"], "Sin turno", None, None])
    return (
        "HORARIOS POR TURNO (SEMANAL)",
        ["NOMBRE", "TURNO", "INICIO", "FIN"],
        filas,
        "horarios_semanales",
    )


def sabado(db: Session, anio: int, mes: int) -> tuple[str, list[str], list[Fila], str]:
    plan = sabados.mes(db, anio, mes)
    filas: list[Fila] = []
    con_sabado: set[str] = set()
    for dia in plan["sabados"]:
        fecha = dia["fecha"].strftime("%d/%m/%Y")
        for t in dia["turnos"]:
            for a in t["auxiliares"]:
                con_sabado.add(str(a["id"]))
                filas.append(
                    [
                        a["nombre_completo"],
                        fecha,
                        NOMBRE_TURNO[t["turno"]],
                        _hora(t["hora_inicio"]),
                        _hora(t["hora_fin"]),
                    ]
                )
    for m in _equipo(db):
        if str(m["id"]) not in con_sabado:
            filas.append([m["nombre_completo"], None, "Libre", None, None])
    return (
        f"HORARIOS TURNO SÁBADO {mes:02d}-{anio}",
        ["NOMBRE", "SÁBADO", "TURNO", "INICIO", "FIN"],
        filas,
        f"sabados_{anio}-{mes:02d}",
    )


def libro_xlsx(titulo: str, cabecera: list[str], filas: list[Fila]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    fino = Side(style="thin", color="9DB8AD")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cabecera))
    ws.cell(row=1, column=1, value=titulo).font = Font(
        bold=True, size=14, color="1E3932"
    )
    for i, h in enumerate(cabecera, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="006241")
        c.alignment = Alignment(horizontal="center")
        c.border = borde
    for f, fila in enumerate(filas, start=3):
        for i, valor in enumerate(fila, start=1):
            c = ws.cell(row=f, column=i, value=valor)
            c.border = borde
            if f % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F4F8F5")
    anchos = [32, 14] + [12] * (len(cabecera) - 2)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = ancho
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def hoja_pdf(titulo: str, cabecera: list[str], filas: list[Fila]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=30,
        rightMargin=30,
        topMargin=30,
        bottomMargin=30,
        title=titulo,
    )
    partes = [Paragraph(titulo, getSampleStyleSheet()["Title"]), Spacer(1, 12)]
    datos = [cabecera] + [["" if v is None else str(v) for v in fila] for fila in filas]
    tabla = Table(datos, colWidths=[220, 90, 90, 70, 70][: len(cabecera)], repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#006241")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9DB8AD")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F4F8F5")],
                ),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    partes.append(tabla)
    doc.build(partes)
    return buf.getvalue()
